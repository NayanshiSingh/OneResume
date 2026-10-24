"""Export Service — renders DOCX files and stores export artifacts privately."""

import hashlib
import io
import json
import logging
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.errors import NotFoundError
from app.models.resume import ExportArtifact, Resume
from app.models.user import User
from app.repositories import ArtifactRepo, ResumeRepo
from app.services import output_check
from app.services.pdf_renderer import render_resume_pdf
from app.services.storage import get_storage

logger = logging.getLogger(__name__)

MEDIA_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def _resume_prefix(user_id: str, resume_id: str) -> str:
    return f"users/{user_id}/resumes/{resume_id}"


def store_artifact(db: Session, user_id: str, resume: Resume, format: str, data: bytes,
                   page_count: int | None = None) -> ExportArtifact:
    """Save a rendered file under a key unique to this user, resume, and format."""
    key = f"{_resume_prefix(user_id, resume.id)}/resume.{format}"
    get_storage().save(key, data)
    return ArtifactRepo.create(
        db, resume_id=resume.id, format=format, storage_key=key,
        checksum=hashlib.sha256(data).hexdigest(), size_bytes=len(data), page_count=page_count,
    )


VERSION_ATTEMPTS = 3


def create_resume_record(db: Session, profile_id: str, jd_id: str | None, job_title: str,
                         draft_id: str, draft_revision: int) -> Resume:
    """Create the resume with the next free version, retrying if a concurrent export took it."""
    for _ in range(VERSION_ATTEMPTS):
        version = ResumeRepo.get_next_version(db, profile_id, job_title)
        try:
            return ResumeRepo.create(db, profile_id=profile_id, jd_id=jd_id, job_title=job_title, version=version,
                                     draft_id=draft_id, draft_revision=draft_revision)
        except IntegrityError:
            db.rollback()
    raise RuntimeError("Could not allocate a resume version")


def render_and_store(db: Session, user: User, resume: Resume, resume_data: dict, page_target: int = 1) -> list[str]:
    """Render, read back, and store each format.

    A format that fails to render, or whose text does not read back as the
    approved content, is left out rather than offered. Problems and page
    overflow are recorded on the resume as warnings.
    """
    renderers = {"pdf": render_resume_pdf, "docx": export_to_docx}
    formats, warnings = [], []
    for format, render in renderers.items():
        label = format.upper()
        try:
            data = render(resume_data)
            check = output_check.verify(format, data, resume_data)
        except Exception:  # renderer boundary: report the missing format instead of failing the export
            logger.exception("%s rendering failed for resume %s", label, resume.id)
            warnings.append(f"The {label} could not be created, so only the other format is available.")
            continue
        if not check.ok:
            logger.error("%s for resume %s failed the read-back check: %s", label, resume.id, check)
            problem = "comes out in the wrong order" if check.out_of_order else f"is missing: {', '.join(check.missing[:3])}"
            warnings.append(f"The {label} was not offered because its text {problem}.")
            continue
        if check.pages and check.pages > page_target:
            warnings.append(f"The {label} is {check.pages} pages; your target was {page_target}.")
        store_artifact(db, user.id, resume, format, data, check.pages)
        formats.append(format)
    resume.warnings = json.dumps(warnings)
    db.commit()
    return formats


def artifact_path(artifact: ExportArtifact) -> Path:
    """Path to a stored file, after checking it still matches what was exported."""
    path = get_storage().path(artifact.storage_key)
    if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != artifact.checksum:
        logger.error("Stored %s for resume %s is missing or changed", artifact.format, artifact.resume_id)
        raise NotFoundError("This file is no longer available. Export the draft again.")
    return path


def delete_resume_files(user_id: str, resume_id: str) -> None:
    get_storage().delete_prefix(_resume_prefix(user_id, resume_id))


def download_filename(resume: Resume, format: str) -> str:
    safe_title = "".join(c if c.isalnum() or c in "-_ " else "" for c in resume.job_title).strip() or "Resume"
    return f"{safe_title.replace(' ', '_')}_v{resume.version}.{format}"


def export_to_docx(resume_data: dict) -> bytes:
    """Render resume data to DOCX bytes."""
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    style.font.name = "Calibri"

    # ── Header ────────────────────────────────────────────────
    pi = resume_data.get("personal_info", {})
    if pi:
        name_para = doc.add_paragraph()
        name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = name_para.add_run(pi.get("full_name", ""))
        run.bold = True
        run.font.size = Pt(18)

        contact_parts = []
        if pi.get("email"):
            contact_parts.append(pi["email"])
        if pi.get("phone_number"):
            contact_parts.append(pi["phone_number"])

        if contact_parts:
            contact = doc.add_paragraph(" | ".join(contact_parts))
            contact.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ── Education ─────────────────────────────────────────────
    if resume_data.get("education"):
        doc.add_heading("Education", level=2)
        for edu in resume_data["education"]:
            text = f"{edu['degree']}"
            if edu.get("field_of_study"):
                text += f" in {edu['field_of_study']}"
            text += f" — {edu['institution']}"
            if edu.get("start_year") or edu.get("end_year"):
                years = f"{edu.get('start_year', '')} - {edu.get('end_year', '')}"
                text += f" ({years})"
            doc.add_paragraph(text)

    # ── Experience ────────────────────────────────────────────
    if resume_data.get("experience"):
        doc.add_heading("Experience", level=2)
        for exp in resume_data["experience"]:
            p = doc.add_paragraph()
            run = p.add_run(exp["title"])
            run.bold = True
            if exp.get("subtitle"):
                p.add_run(f"  —  {exp['subtitle']}")
            for bullet in exp.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")

    # ── Projects ──────────────────────────────────────────────
    if resume_data.get("projects"):
        doc.add_heading("Projects", level=2)
        for proj in resume_data["projects"]:
            p = doc.add_paragraph()
            run = p.add_run(proj["title"])
            run.bold = True
            if proj.get("subtitle"):
                p.add_run(f"  |  {proj['subtitle']}")
            for bullet in proj.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")

    # ── Skills ────────────────────────────────────────────────
    if resume_data.get("skills"):
        doc.add_heading("Skills", level=2)
        doc.add_paragraph(", ".join(resume_data["skills"]))

    # ── Certifications ────────────────────────────────────────
    if resume_data.get("certifications"):
        doc.add_heading("Certifications", level=2)
        for cert in resume_data["certifications"]:
            text = cert["name"]
            if cert.get("issuing_organization"):
                text += f" — {cert['issuing_organization']}"
            if cert.get("year"):
                text += f" ({cert['year']})"
            doc.add_paragraph(text)

    # ── Achievements ──────────────────────────────────────────
    if resume_data.get("achievements"):
        doc.add_heading("Achievements", level=2)
        for ach in resume_data["achievements"]:
            text = ach["title"]
            if ach.get("description"):
                text += f": {ach['description']}"
            doc.add_paragraph(text, style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
